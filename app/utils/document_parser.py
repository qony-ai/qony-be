from __future__ import annotations

import asyncio
from pathlib import Path


async def parse_document(file_path: str) -> str:
    return await asyncio.to_thread(_parse_document_sync, file_path)


def _parse_document_sync(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if suffix == ".docx":
        from docx import Document

        document = Document(file_path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()).strip()

    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(file_path)
        lines: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if isinstance(text, str) and text.strip():
                    lines.append(text.strip())
        return "\n".join(lines).strip()

    return path.read_text(encoding="utf-8", errors="ignore").strip()
